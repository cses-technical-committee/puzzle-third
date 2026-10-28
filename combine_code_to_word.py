from pathlib import Path
import tkinter as tk
from tkinter import filedialog
from docx import Document
from docx.shared import Pt, Inches
import re

SKIP_DIRS = {
    "node_modules", ".git", "dist", "build",
    "__pycache__", ".venv", "venv", "coverage"
}

EXTENSIONS = {
    ".js", ".jsx", ".ts", ".tsx", ".mjs",
    ".html", ".htm", ".css", ".scss", ".sass",
    ".py", ".java", ".c", ".cpp", ".h", ".cs",
    ".php", ".json", ".xml", ".yaml", ".yml",
    ".md", ".txt", ".sql", ".sh", ".bat",
    ".vue", ".svelte", ".astro", ".rb", ".go",
    ".rs", ".kt", ".swift", ".dart", ".r",
    ".toml", ".ini", ".conf", ".graphql"
}

SPECIAL_FILES = {
    "dockerfile", "makefile", ".gitignore",
    ".dockerignore", ".editorconfig", "readme"
}


def clean_for_word(text):
    """Remove XML-invalid characters before adding text to Word."""
    return re.sub(
        r"[\x00-\x08\x0B\x0C\x0E-\x1F\uFFFE\uFFFF]",
        "",
        text
    )


def main():
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    try:
        folder = filedialog.askdirectory(
            title="Select your project folder"
        )
    finally:
        root.destroy()

    if not folder:
        print("No folder selected.")
        return

    project = Path(folder)
    output = project / "Code_Combined.docx"
    files = []

    for file in project.rglob("*"):
        try:
            if not file.is_file():
                continue

            relative = file.relative_to(project)

            if any(
                part.lower() in SKIP_DIRS
                for part in relative.parts[:-1]
            ):
                continue

            if file.name == output.name:
                continue

            if (
                file.suffix.lower() in EXTENSIONS
                or file.name.lower() in SPECIAL_FILES
            ):
                if file.stat().st_size <= 2 * 1024 * 1024:
                    files.append(file)

        except OSError as error:
            print(f"Skipping {file}: {error}")

    files.sort(
        key=lambda f: str(f.relative_to(project)).lower()
    )

    if not files:
        print("No supported code files found.")
        return

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    doc.add_heading("Combined Project Code", 0)
    doc.add_paragraph(f"Project folder: {project}")
    doc.add_paragraph(f"Total files: {len(files)}")
    doc.add_paragraph(
        "Each section contains the relative file path "
        "and the corresponding source code."
    )
    doc.add_page_break()

    for index, file in enumerate(files, start=1):
        relative_name = file.relative_to(project).as_posix()

        doc.add_heading(
            f"{index}. {clean_for_word(relative_name)}",
            level=1
        )

        try:
            code = file.read_text(
                encoding="utf-8-sig",
                errors="replace"
            )
            code = clean_for_word(code)

        except Exception as error:
            doc.add_paragraph(
                f"Could not read file: {clean_for_word(str(error))}"
            )
            continue

        if not code.strip():
            doc.add_paragraph("[Empty file]")
        else:
            for line in code.splitlines():
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)

                run = paragraph.add_run(line if line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(8)

        if index < len(files):
            doc.add_page_break()

    try:
        doc.save(output)
    except Exception as error:
        print(f"Could not save Word document: {error}")
        return

    print(f"\nSuccessfully combined {len(files)} files!")
    print(f"Word document created at:\n{output}")


if __name__ == "__main__":
    main()